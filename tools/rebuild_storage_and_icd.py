from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Docs" / "ProjectsDocs"


def setup(doc, footer_text):
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.footer.paragraphs[0].alignment = 1
        section.footer.paragraphs[0].text = footer_text


def title(doc, lines):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = 1
        r = p.add_run(line)
        r.bold = i == 0
        r.font.size = Pt(16 if i == 0 else 12)


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def p(doc, text):
    doc.add_paragraph(text)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = str(value)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    doc.add_paragraph()


def code(doc, text):
    q = doc.add_paragraph()
    q.style = "No Spacing"
    r = q.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(8)


def build_data_model():
    d = Document()
    title(d, [
        "DATA MODEL & LOCAL STORAGE SCHEMA",
        "Clinker Bucket Elevator Vision Monitoring System",
        "SQLite alerts + local server evidence",
        "Document Status: Draft v0.3 — revised to match the physical-server MVP",
    ])
    table(d, ["Version", "Change"], [["0.3", "Replaced VM/NAS/outbox/mobile-backend assumptions with local SQLite, local HDD evidence, and the server-local web/API viewer." ]])

    h(d, "1 Introduction")
    h(d, "1.1 Purpose", 2)
    p(d, "This document defines the data stored by the MVP pipeline on the physical Ubuntu server. It covers the SQLite alert database, full-frame evidence JPEGs, local REST read behavior, timestamps, retention, and validation rules. It does not define model internals or a future mobile application.")
    h(d, "1.2 MVP storage decisions", 2)
    bullets(d, [
        "SQLite is stored on the server's local 500 GB HDD.",
        "Evidence JPEGs are stored on the same local disk; NAS/NFS is not used in the MVP.",
        "Only qualifying detections are stored; normal empty model results and unprocessed frames are not stored.",
        "One alert row is created per qualifying detection. No deduplication or cooldown is applied.",
        "The database is read by the local REST/web viewer in the pipeline container.",
        "There is no external alert backend, delivery state, push notification, mobile feed, or operator account in the MVP.",
        "The target retention period is 365 days, subject to local disk capacity and cleanup success.",
    ])

    h(d, "2 Storage layout")
    table(d, ["Data", "Path", "Notes"], [
        ["SQLite", "/srv/clinker-vision/data/alerts.db", "SQLite WAL mode; pipeline writer and local API reader."],
        ["Evidence", "/srv/clinker-vision/evidence/YYYY/MM/DD/CAM-1/<alert_id>.jpg", "Full-frame JPEG with all detections overlaid."],
        ["Configuration", "/srv/clinker-vision/config/config.yaml", "Protected permissions; contains NVR connection values."],
        ["Logs", "/srv/clinker-vision/logs/ or Docker logs", "Must be rotated and must not contain secrets or images."],
    ])

    h(d, "3 Alert schema")
    p(d, "The following is the normative logical schema. The exact migration mechanism may be chosen during implementation.")
    code(d, """CREATE TABLE IF NOT EXISTS alerts (
  alert_id TEXT PRIMARY KEY,
  captured_at TEXT NOT NULL,
  detected_at TEXT NOT NULL,
  processed_at TEXT,
  created_at TEXT NOT NULL,
  camera_id TEXT NOT NULL CHECK (camera_id IN ('CAM-1','CAM-2','CAM-3','CAM-4','CAM-5','CAM-6')),
  observation_target TEXT NOT NULL CHECK (observation_target IN ('godet','galet','clinker_level')),
  fault_type TEXT NOT NULL,
  confidence REAL NOT NULL CHECK (confidence >= 0 AND confidence <= 1),
  frame_id TEXT NOT NULL,
  model_version TEXT NOT NULL,
  bbox_x REAL,
  bbox_y REAL,
  bbox_w REAL,
  bbox_h REAL,
  scalar_value REAL,
  evidence_ref TEXT NOT NULL UNIQUE,
  seen_at TEXT
);

CREATE INDEX IF NOT EXISTS idx_alerts_detected
  ON alerts(detected_at DESC);
CREATE INDEX IF NOT EXISTS idx_alerts_camera_detected
  ON alerts(camera_id, detected_at DESC);
CREATE INDEX IF NOT EXISTS idx_alerts_seen_detected
  ON alerts(seen_at, detected_at DESC);""")
    table(d, ["Field", "Rule"], [
        ["alert_id", "UUID v4 generated by the pipeline; never reused."],
        ["captured_at", "Timestamp attached to the source frame."],
        ["detected_at", "Detection event time; normally captured_at for a frame-based alert."],
        ["processed_at", "Model response time, when supplied."],
        ["created_at", "SQLite row creation time."],
        ["camera_id", "CAM-1 in MVP; future rows may use CAM-2 through CAM-6."],
        ["observation_target", "Value returned per detection by the model; godet in MVP."],
        ["fault_type", "Model-reported fault; CRACK or MISALIGNMENT are the intended MVP values."],
        ["confidence", "Raw model confidence between 0 and 1; pipeline threshold begins at 0.85."],
        ["frame_id/model_version", "Traceability to the exact inference request and model build."],
        ["bbox_*", "Normalized full-frame top-left coordinates. Required for boxed detections."],
        ["scalar_value", "Reserved for future scalar detections; null for MVP godet faults."],
        ["evidence_ref", "Relative path from /srv/clinker-vision/evidence, not a NAS path or URL."],
        ["seen_at", "Null until the local viewer marks an alert seen; one shared timestamp, no user identity."],
    ])
    p(d, "For MVP alerts, bbox_x, bbox_y, bbox_w, and bbox_h must all be present and within 0.0–1.0. The pipeline validates and clamps model coordinates before rendering and storing them. A future scalar-only alert may leave the bounding box null and populate scalar_value.")

    h(d, "4 Evidence specification")
    table(d, ["Property", "MVP rule"], [
        ["Source", "The full JPEG frame sent to the model; no pipeline crop."],
        ["Overlay", "All model detections from the frame, red boxes, and fault/confidence labels."],
        ["Encoding", "JPEG quality 85 initially; no EXIF orientation dependency."],
        ["Path", "/srv/clinker-vision/evidence/YYYY/MM/DD/<camera_id>/<alert_id>.jpg"],
        ["Write", "Write to .tmp in the same directory, fsync/close as appropriate, then rename atomically."],
        ["Failure", "If the evidence write fails, log the error and do not present the alert as complete."],
    ])
    p(d, "The local viewer serves evidence only through the pipeline's local API. No remote file server, signed URL, NAS mount, or mobile evidence URL exists in the MVP.")

    h(d, "5 Write and read behavior")
    numbered(d, [
        "Receive a valid model detection that passes the configured threshold.",
        "Validate target, fault, confidence, and full-frame bounding box.",
        "Render all frame detections onto the full source frame.",
        "Write the evidence JPEG atomically to the local evidence path.",
        "Insert the alert row into SQLite in a short transaction.",
        "Make the row available through the local REST API and web viewer.",
    ])
    p(d, "If the process stops between the file write and database insert, startup cleanup may remove orphan JPEGs. If it stops before the file rename, the temporary file may be removed by the next cleanup pass. There is no persistent frame queue or alert-delivery outbox in this MVP.")

    h(d, "6 Local read model")
    table(d, ["Operation", "Behavior"], [
        ["List", "GET /api/v1/alerts, sorted by detected_at descending."],
        ["Filters", "camera_id, fault_type, unseen_only, since, limit, and cursor may be supported."],
        ["Detail", "GET /api/v1/alerts/{alert_id} returns metadata and evidence path/URL for the local viewer."],
        ["Evidence", "GET /api/v1/alerts/{alert_id}/evidence serves the local JPEG."],
        ["Seen", "POST /api/v1/alerts/{alert_id}/seen sets seen_at if null; repeated calls are no-ops."],
        ["Access", "Local server only; no authentication in MVP."],
    ])

    h(d, "7 Retention and capacity")
    bullets(d, [
        "Target retention is 365 days for alert rows and evidence JPEGs.",
        "Cleanup runs daily at a configured time and deletes data older than 365 days.",
        "Cleanup must never follow paths outside the configured evidence root.",
        "The service must warn at a configurable low-disk threshold and stop safely if it cannot write evidence.",
        "The 500 GB disk is expected to provide comfortable MVP capacity, but real alert rate and JPEG size must be measured.",
        "No independent database/evidence backup is included in the MVP; this is an accepted project risk.",
    ])

    h(d, "8 Time and traceability rules")
    bullets(d, [
        "Store plant-local ISO 8601 timestamps with numeric offset; the current project timezone is Africa/Casablanca.",
        "Preserve captured_at from the frame; never replace it with send time.",
        "Use frame_id, camera_id, sequence_no in logs to diagnose drops and reordering.",
        "Keep model_version on every alert so an evidence image can be traced to the model build.",
        "If wall-clock ordering becomes ambiguous during a timezone transition, use created_at and alert_id as tie-breakers.",
    ])

    h(d, "9 Open items")
    table(d, ["ID", "Item", "Owner"], [
        ["DM-OI-1", "Confirm JPEG size and disk growth from real CAM-1 frames.", "Project owner"],
        ["DM-OI-2", "Confirm exact REST JSON field names and pagination implementation.", "Developer"],
        ["DM-OI-3", "Confirm cleanup time and low-disk threshold.", "Project owner"],
        ["DM-OI-4", "Confirm that no backup is required during the MVP test period.", "Project owner"],
    ])
    save = OUT / "clinker-vision-datamodel.docx"
    setup(d, "Data Model v0.3 — Draft")
    d.save(save)


def build_icd():
    d = Document()
    title(d, [
        "INTERFACE CONTROL DOCUMENT",
        "Clinker Bucket Elevator Vision Monitoring System",
        "Model gRPC and local web/API boundaries",
        "Document Status: Draft v0.4 — revised to match the physical-server MVP",
    ])
    table(d, ["Version", "Change"], [["0.4", "Removed the obsolete mobile/backend notification boundary from the MVP; added the local REST/web boundary and local Docker transport." ]])

    h(d, "1 Introduction")
    h(d, "1.1 Purpose", 2)
    p(d, "This document defines the interfaces crossing the pipeline boundary in the MVP: the local gRPC model-service interface and the local REST/web read interface. The model's internal implementation and the web UI's internal implementation are out of scope.")
    h(d, "1.2 MVP ownership", 2)
    bullets(d, [
        "The pipeline team owns the Go client, local REST API, local web viewer, and data persistence.",
        "The model owner supplies the Python/OpenCV model container and implements the gRPC server.",
        "A future mobile/backend contract is deferred and is not an MVP dependency.",
    ])

    h(d, "2 Interface 1 — Model service")
    h(d, "2.1 Transport", 2)
    p(d, "The pipeline and model run as separate Docker containers on the same physical Ubuntu server. They communicate over a private Docker Compose network using gRPC streaming. TLS is not used in the MVP because the model port is not published to the host. One persistent stream is used per inference worker, and responses are correlated by frame_id rather than assumed to be ordered.")
    h(d, "2.2 Package and versioning", 2)
    p(d, "The proposed proto package is clinkervision.inference.v1. Incompatible schema changes require a package-version bump and review by both owners. model_version identifies model weights/build and is separate from the proto package version.")
    h(d, "2.3 InferenceRequest", 2)
    table(d, ["Field", "Type", "Rule"], [
        ["frame_id", "string UUID", "Generated by pipeline; echoed in response."],
        ["camera_id", "string", "CAM-1 in MVP; future CAM-2 through CAM-6."],
        ["image_data", "bytes", "Full-frame JPEG; pipeline performs no crop."],
        ["captured_at", "Timestamp", "Original frame capture time."],
        ["sequence_no", "uint64", "Monotonic per camera_id; useful for drop/reorder diagnosis."],
    ])
    code(d, """message InferenceRequest {
  string frame_id = 1;
  string camera_id = 2;
  bytes image_data = 3; // full-frame JPEG
  google.protobuf.Timestamp captured_at = 4;
  uint64 sequence_no = 5; // per camera_id
}""")
    h(d, "2.4 InferenceResponse and Detection", 2)
    table(d, ["Field", "Type", "Rule"], [
        ["frame_id", "string UUID", "Must echo the request frame_id."],
        ["detections", "repeated Detection", "Empty means the frame was processed and no qualifying finding was returned."],
        ["scalar_measurements", "map<string,float>", "Reserved for future scalar targets; not required for godet MVP."],
        ["model_version", "string", "Required on every response."],
        ["processed_at", "Timestamp", "Model response time."],
        ["observation_target", "enum", "GODET for MVP; value is per detection."],
        ["fault_type", "enum", "CRACK or MISALIGNMENT are intended MVP values."],
        ["confidence", "float", "Raw 0.0–1.0; pipeline applies threshold."],
        ["bounding_box", "BoundingBox", "Normalized 0.0–1.0, top-left, relative to full frame."],
    ])
    code(d, """message InferenceResponse {
  string frame_id = 1;
  repeated Detection detections = 2;
  map<string, float> scalar_measurements = 3;
  string model_version = 4;
  google.protobuf.Timestamp processed_at = 5;
}

message Detection {
  ObservationTarget observation_target = 1;
  FaultType fault_type = 2;
  float confidence = 3;
  BoundingBox bounding_box = 4;
}""")
    h(d, "2.5 Model responsibilities", 2)
    bullets(d, [
        "Decode the full JPEG and perform all dynamic crop, resize, and normalization in Python/OpenCV.",
        "Use the same preprocessing path in training and serving.",
        "Return raw confidence; do not apply the pipeline's 0.85 threshold.",
        "Return model_version on every response.",
        "Return full-frame coordinates even when inference uses an internal crop.",
    ])
    h(d, "2.6 Errors", 2)
    table(d, ["Condition", "Pipeline behavior"], [
        ["OK with empty detections", "Normal result; no alert."],
        ["UNAVAILABLE", "Log model failure; discard unprocessed frame; dependency timer starts."],
        ["DEADLINE_EXCEEDED", "Log timeout and latency; discard frame; dependency timer starts."],
        ["INVALID_ARGUMENT", "Log frame_id/camera_id; discard malformed frame; do not retry the same frame."],
        ["Invalid response", "Log and discard response; validate before alert creation."],
        ["Required dependency unavailable for five minutes", "Pipeline stops completely with a clear reason."],
    ])

    h(d, "3 Interface 2 — Local REST/web viewer")
    h(d, "3.1 Transport and access", 2)
    p(d, "The pipeline container serves HTTP on 127.0.0.1 only. The operator opens the viewer directly on the Ubuntu server. There is no authentication, TLS, remote access, external backend, bearer token, mobile app, or push service in the MVP.")
    h(d, "3.2 List alerts", 2)
    code(d, "GET /api/v1/alerts?camera_id=CAM-1&fault_type=CRACK&unseen_only=true&limit=50&cursor=<opaque>")
    p(d, "The response is JSON containing alert metadata and a local evidence URL. Results are sorted by detected_at descending. Pagination may use an opaque cursor; a simple limit-based implementation is acceptable for the first local MVP if documented.")
    h(d, "3.3 Read one alert and evidence", 2)
    code(d, "GET /api/v1/alerts/{alert_id}")
    code(d, "GET /api/v1/alerts/{alert_id}/evidence")
    p(d, "The evidence endpoint serves the JPEG from the configured local evidence root after validating that the requested path stays inside that root. It must not expose arbitrary filesystem paths.")
    h(d, "3.4 Mark seen", 2)
    code(d, "POST /api/v1/alerts/{alert_id}/seen")
    code(d, "200 { \"alert_id\": \"...\", \"seen_at\": \"2026-09-17T14:05:01+01:00\" }")
    p(d, "The first call sets the shared seen_at value. Repeated calls are idempotent. There are no users or per-device seen states.")

    h(d, "4 Evidence contract")
    table(d, ["Field", "MVP rule"], [
        ["evidence_ref", "Relative path below /srv/clinker-vision/evidence; never a NAS path or remote URL."],
        ["Image", "Full frame sent to model, with all detections overlaid."],
        ["Coordinates", "Normalized full-frame top-left origin."],
        ["JPEG", "Quality 85 initially; expected size must be measured on real frames."],
        ["Write", "Temporary file and atomic rename."],
    ])

    h(d, "5 Change management and open items")
    bullets(d, [
        "The model owner must sign off the proto, full-frame coordinate convention, input size, preprocessing parity, and CPU benchmark.",
        "The developer must sign off the local REST JSON shape and web port before implementation is considered complete.",
        "Any remote/mobile interface requires a new contract, TLS, authentication, and security review.",
        "The exact NVR details are deployment configuration, not part of the model proto.",
    ])
    table(d, ["ID", "Open item", "Owner"], [
        ["ICD-OI-1", "Obtain final proto and model Docker image.", "Model owner"],
        ["ICD-OI-2", "Confirm model CPU/RAM requirements and benchmark.", "Model owner"],
        ["ICD-OI-3", "Confirm local REST JSON and web port.", "Developer/project owner"],
        ["ICD-OI-4", "Confirm exact NVR/CAM-1 RTSP input during deployment.", "Project owner"],
    ])
    save = OUT / "clinker-vision-icd.docx"
    setup(d, "ICD v0.4 — Draft")
    d.save(save)


if __name__ == "__main__":
    build_data_model()
    build_icd()
    print("Rebuilt data model and ICD")
